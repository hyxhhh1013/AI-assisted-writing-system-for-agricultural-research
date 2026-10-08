"""Fused stage-report deck — denser layout (navy / rust, not the old forest-gold theme).

Usage: python scripts/build-stage-report-pptx.py
Output: docs/reports/阶段汇报-2026-09-30-下一期投入.pptx
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports" / "阶段汇报-2026-09-30-下一期投入.pptx"

FONT = "Microsoft YaHei"
NAVY = RGBColor(0x14, 0x2A, 0x45)
NAVY_SOFT = RGBColor(0x1E, 0x3D, 0x5C)
RUST = RGBColor(0xC0, 0x56, 0x2A)
SAND = RGBColor(0xF3, 0xEC, 0xE0)
CREAM = RGBColor(0xFB, 0xF8, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x27, 0x34)
MUTED = RGBColor(0x5B, 0x67, 0x75)
LINE = RGBColor(0xD9, 0xD0, 0xC2)
PALE = RGBColor(0xEA, 0xF0, 0xF6)
ROW = RGBColor(0xF7, 0xF2, 0xE9)

SLIDE_W = 13.333
SLIDE_H = 7.5
TOTAL = 22


def set_font(run, size: float, bold: bool, color: RGBColor) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = r_pr.find(qn(f"a:{tag}"))
        if el is None:
            el = r_pr.makeelement(qn(f"a:{tag}"), {})
            r_pr.append(el)
        el.set("typeface", FONT)


def shape(slide, kind, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, lw=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    return s


def rect(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def txt(
    slide,
    text: str,
    x,
    y,
    w,
    h,
    size=16,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set(
        "anchor",
        {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
    )
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def lines(slide, items: list[str], x, y, w, h, size=15, color=INK, gap=7, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = item
        set_font(run, size, bold, color)
    return box


def border_cell(cell, color="D9D0C2", sz="5080") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        existing = tc_pr.find(qn(f"a:{edge}"))
        if existing is not None:
            tc_pr.remove(existing)
        ln = etree.SubElement(tc_pr, qn(f"a:{edge}"))
        ln.set("w", sz)
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
        srgb.set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def paint_cell(cell, text, size, bold, color, fill) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.12)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.08)
    cell.margin_bottom = Inches(0.08)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    border_cell(cell)


def table(slide, rows: list[list[str]], x, y, w, h, col_w: list[float], size=13):
    sh = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = sh.table
    for i, width in enumerate(col_w):
        tbl.columns[i].width = Inches(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            header = r == 0
            fill = NAVY if header else (WHITE if r % 2 else ROW)
            color = WHITE if header else INK
            paint_cell(tbl.cell(r, c), text, size, header or c == 0, color, fill)
    return tbl


def note(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def page(prs, n: int, eyebrow: str, title: str, spoken: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    rect(slide, 0, 0, 0.18, SLIDE_H, NAVY)
    rect(slide, 0.18, 0, SLIDE_W - 0.18, 1.18, NAVY)
    txt(slide, f"{n:02d}  /  {TOTAL}     {eyebrow}", 0.5, 0.14, 12.3, 0.28, 11, True, RUST)
    txt(slide, title, 0.48, 0.42, 12.4, 0.58, 24, True, WHITE)
    rect(slide, 0.18, 7.18, SLIDE_W - 0.18, 0.32, SAND)
    txt(slide, "禾书耕文  ·  实验室科研写作  ·  阶段汇报", 0.5, 7.2, 8.5, 0.26, 11, False, MUTED)
    txt(slide, f"{n} / {TOTAL}", 11.2, 7.2, 1.7, 0.26, 11, False, MUTED, PP_ALIGN.RIGHT)
    note(slide, spoken)
    return slide


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "禾书耕文：阶段汇报与下一期投入"
    prs.core_properties.subject = "2026-10 融合稿"
    prs.core_properties.author = "禾书耕文"
    for fn in (
        cover,
        agenda,
        judgment,
        market,
        price,
        wedge,
        positioning,
        phases,
        foundations,
        plagiarism,
        library,
        charts,
        agent,
        path,
        errors,
        case,
        artifacts,
        boundary,
        next_phase,
        difficulty,
        support,
        close,
    ):
        fn(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


def cover(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 5.15, SLIDE_H, NAVY)
    rect(slide, 5.15, 0, 8.183, SLIDE_H, CREAM)
    rect(slide, 5.15, 0, 0.08, SLIDE_H, RUST)
    txt(slide, "阶段汇报", 0.45, 1.35, 4.4, 0.32, 13, True, RUST)
    txt(slide, "禾书耕文", 0.42, 1.85, 4.5, 0.95, 40, True, WHITE)
    txt(slide, "实验室科研写作系统", 0.45, 2.85, 4.4, 0.45, 16, False, RGBColor(0xC5, 0xD0, 0xDC))
    txt(slide, "2026.10\n课题组导师 / 中心", 0.45, 5.85, 4.4, 0.85, 14, False, RGBColor(0x9A, 0xAE, 0xC0))
    txt(slide, "今天要定的事", 5.65, 1.45, 7.1, 0.32, 13, True, RUST)
    txt(
        slide,
        "请老师给使用闭环：\n定性、一个人、一道真题、一个看稿日。",
        5.65,
        1.9,
        7.1,
        1.45,
        24,
        True,
        NAVY,
    )
    blocks = [
        ("系统层", "查重、组内文献、配图、Agent 已串成一条初稿路径。"),
        ("稿件层", "龙智、陈韶光、易程三篇真实题目初稿，用来看 AI 上限。"),
        ("不承诺", "完整产品、一键全文。优势是人在组里一直改。"),
    ]
    for i, (h, d) in enumerate(blocks):
        y = 3.55 + i * 0.95
        rect(slide, 5.65, y, 7.05, 0.85, WHITE, LINE)
        rect(slide, 5.65, y, 0.1, 0.85, RUST if i == 0 else NAVY)
        txt(slide, h, 5.95, y + 0.08, 6.5, 0.28, 13, True, NAVY)
        txt(slide, d, 5.95, y + 0.38, 6.5, 0.38, 13, False, MUTED)
    note(slide, "封面先把请求说死。框架能出初稿，三篇学长稿标定上限，缺的是作者自己走一遍。")


def agenda(prs) -> None:
    slide = page(prs, 2, "目录", "十五分钟，六件事", "时间不够留 03、05、07、08、09、16、21、22。")
    items = [
        ("01", "市场", "课题组已经在为什么付钱", "和公网订阅比量级，不和品牌比功能个数。"),
        ("02", "定位", "占哪条缝，三期怎样算过", "驻组、作者主笔、可审校初稿。过不了就停。"),
        ("03", "基础", "查重、知识库、图表、Agent", "四件原来就要花钱或花时间的事，收进一条路径。"),
        ("04", "证明", "一条路径，三篇学长原稿", "代跑标定上限，不是组员已经独立用起来。"),
        ("05", "难点", "技术能扛，现实层在空转", "定性没锁，组里暂时也不急着写论文。"),
        ("06", "支持", "定性、人、题、看稿日", "不请大额开发费，不请第二名全职开发。"),
    ]
    for i, (num, name, lead, rest) in enumerate(items):
        col, row = i % 2, i // 2
        x, y = 0.48 + col * 6.35, 1.42 + row * 1.8
        rect(slide, x, y, 6.15, 1.62, WHITE, LINE)
        rect(slide, x, y, 1.15, 1.62, NAVY if i % 2 == 0 else NAVY_SOFT)
        txt(slide, num, x + 0.08, y + 0.5, 1.0, 0.55, 20, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, name, x + 1.35, y + 0.18, 4.55, 0.38, 18, True, NAVY)
        txt(slide, lead, x + 1.35, y + 0.58, 4.55, 0.35, 13, True, RUST)
        txt(slide, rest, x + 1.35, y + 0.98, 4.55, 0.42, 13, False, MUTED)


def judgment(prs) -> None:
    slide = page(prs, 3, "阶段判断", "框架能出初稿，上限已经看见", "还没证明的是组员自己打开系统走完。")
    steps = ["方向规划", "建项目", "分节写作", "审查 · 查重 · 配图", "Word / PDF"]
    gap = 0.12
    box_w = (12.35 - gap * 4) / 5
    for i, step in enumerate(steps):
        x = 0.48 + i * (box_w + gap)
        fill = NAVY if i == 4 else WHITE
        fg = WHITE if i == 4 else NAVY
        rect(slide, x, 1.4, box_w, 0.95, fill, None if i == 4 else LINE)
        txt(slide, step, x + 0.06, 1.52, box_w - 0.12, 0.7, 13, True, fg, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < 4:
            txt(slide, "→", x + box_w - 0.02, 1.62, 0.22, 0.5, 14, True, RUST, PP_ALIGN.CENTER)
    cards = [
        ("安全", "稿不丢，页面和上传收紧，工作台以后好改。"),
        ("真题", "指定人、指定题、指定截止日期。走出作者自己点出来的 1～2 篇。"),
        ("尺子", "定性锁在驻组、可审校初稿。按问题清单改小功能，事先说上限。"),
    ]
    for i, (h, d) in enumerate(cards):
        x = 0.48 + i * 4.2
        rect(slide, x, 2.6, 4.02, 2.55, WHITE, LINE)
        rect(slide, x, 2.6, 4.02, 0.12, RUST if i == 1 else NAVY)
        txt(slide, f"0{i + 1}", x + 0.22, 2.88, 3.55, 0.32, 12, True, RUST)
        txt(slide, h, x + 0.22, 3.25, 3.55, 0.45, 20, True, NAVY)
        txt(slide, d, x + 0.22, 3.8, 3.55, 1.1, 14, False, INK)
    txt(slide, "仓库约四百余次提交。近期主轴是每轮一个可见结果，蓝图和大纲都等人批准。", 0.48, 5.35, 12.35, 0.55, 13, False, MUTED)


def market(prs) -> None:
    slide = page(prs, 4, "市场", "四类工具，各卖什么", "先承认别人强。我们只讲实验室这一摊。")
    table(
        slide,
        [
            ["类型", "代表", "课题组用它解决什么"],
            ["公网润色 / 投稿检查", "Paperpal", "英文语法、投稿前语言和引用检查"],
            ["公网文献助手", "SciSpace", "海量论文里检索、和 PDF 对话、写综述笔记"],
            ["综述 Agent", "Elicit", "筛文献、抽表格、做系统综述"],
            ["通用对话", "ChatGPT 等", "零门槛起草、改写、答疑"],
        ],
        0.48,
        1.4,
        12.35,
        4.35,
        [3.5, 2.4, 6.45],
        14,
    )
    rect(slide, 0.48, 5.95, 12.35, 0.95, SAND)
    txt(
        slide,
        "国内「科研全家桶」功能更多、培训更强。不拿功能个数去比，避免被问成「怎么打过解螺旋」。我们只讲组内文献和实验数据进稿。",
        0.7,
        6.1,
        11.95,
        0.65,
        14,
        False,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def price(prs) -> None:
    slide = page(prs, 5, "价格锚", "和课题组已经在花的钱比", "公开标价说大约。我们要的是同一量级的服务器、接口和陪跑。")
    triples = [
        ("Paperpal", "144～348 美元 / 人 / 年", "Prime 到 Pro，英文润色与投稿检查"),
        ("SciSpace", "年付约 12 美元 / 月", "Premium 约 144 美元 / 年，文献任务重时更贵"),
        ("Elicit Pro", "约 49 美元 / 人 / 月", "按年付，偏系统综述，比写作订阅贵"),
    ]
    for i, (a, b, c) in enumerate(triples):
        x = 0.48 + i * 4.2
        rect(slide, x, 1.4, 4.02, 2.35, WHITE, LINE)
        rect(slide, x, 1.4, 4.02, 0.1, RUST if i == 0 else NAVY)
        txt(slide, a, x + 0.22, 1.65, 3.55, 0.35, 14, True, RUST)
        txt(slide, b, x + 0.22, 2.1, 3.55, 0.7, 16, True, NAVY)
        txt(slide, c, x + 0.22, 2.85, 3.55, 0.65, 13, False, MUTED)
    rect(slide, 0.48, 3.95, 7.7, 2.9, WHITE, LINE)
    txt(slide, "查重是另一笔反复开支", 0.72, 4.15, 7.2, 0.35, 15, True, NAVY)
    lines(
        slide,
        [
            "知网个人约 1.5 元/千字，AIGC 约 2 元/千字。",
            "硕士稿三万字一轮大约一百出头，改一版再测就是几百。",
            "终检不能省。能省的是还没改完就反复送正式库。",
        ],
        0.72,
        4.65,
        7.2,
        1.9,
        14,
    )
    rect(slide, 8.38, 3.95, 4.45, 2.9, NAVY)
    txt(slide, "我们请的量级", 8.62, 4.2, 4.05, 0.35, 13, True, RUST)
    txt(
        slide,
        "五人组一年公网订阅大约几千到一万出头人民币。文献和数据在对方云上。下一期买的不是完整产品，是把组内路径绑到真实写稿任务上。",
        8.62,
        4.7,
        4.05,
        1.85,
        14,
        False,
        WHITE,
    )


def wedge(prs) -> None:
    slide = page(prs, 6, "差异", "我们占的那条缝", "公网订阅继续润色英文、查海量文献。禾书耕文补进不了组内数据的那一段。")
    points = [
        ("01", "资产留在组里", "组内 PDF、实验表、稿留在实验室服务器上。"),
        ("02", "引用有边界", "写作落到当前这一节，引用编号不得超出本次检索到的文献。"),
        ("03", "数字要对上", "结果章要接得上实验数据，不能凭对话编出数字。"),
        ("04", "农科写法在流程里", "IMRaD、结果与讨论分开，过度绝对的措辞会拦。"),
    ]
    for i, (n, h, d) in enumerate(points):
        y = 1.4 + i * 1.28
        rect(slide, 0.48, y, 12.35, 1.16, WHITE, LINE)
        rect(slide, 0.48, y, 1.35, 1.16, NAVY)
        txt(slide, n, 0.48, y + 0.32, 1.35, 0.5, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, h, 2.05, y + 0.16, 10.4, 0.38, 17, True, NAVY)
        txt(slide, d, 2.05, y + 0.58, 10.4, 0.42, 14, False, MUTED)


def positioning(prs) -> None:
    slide = page(prs, 7, "定位", "作者主笔。优势是驻组，不是完整产品", "听到「一键出全文」回到这页。")
    rect(slide, 0.48, 1.4, 12.35, 1.85, NAVY)
    txt(
        slide,
        "我不保证做出一个完整的、别人拿去就能用的产品。我保证的是：人一直在课题组里，按真实写稿需求改，做成最贴这组的系统。",
        0.75,
        1.6,
        11.85,
        1.45,
        18,
        True,
        WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    duties = [("检索", "从组内文献里找依据"), ("起草", "按当前这一节出草稿"), ("把关", "配图、审查、查重预检"), ("交稿", "导出给导师看的 Word / PDF")]
    for i, (h, d) in enumerate(duties):
        x = 0.48 + i * 3.15
        rect(slide, x, 3.45, 3.02, 1.55, WHITE, LINE)
        txt(slide, h, x + 0.15, 3.6, 2.72, 0.4, 16, True, RUST, PP_ALIGN.CENTER)
        txt(slide, d, x + 0.15, 4.1, 2.72, 0.65, 13, False, INK, PP_ALIGN.CENTER)
    txt(
        slide,
        "学生按阶段往下走。大纲、蓝图、每一轮结果，作者点头之后才进入下一步。不是外包交钥匙，也不是拿功能个数去打公网品牌。",
        0.48,
        5.2,
        12.35,
        0.85,
        14,
        False,
        MUTED,
    )


def phases(prs) -> None:
    slide = page(prs, 8, "投资", "三期分开验收，过不了就停", "现在接近 A 的系统能走完。缺的是人真的自己走完。")
    table(
        slide,
        [
            ["期", "这期买到什么", "怎样算过"],
            ["A  本实验室打穿", "算力 + 陪跑", "系统能走完；三篇代跑可看上限；另有 1 篇作者自己点出来，导师看过"],
            ["B  组内能复制", "按清单小改，有上限", "至少 3 人独立完成「文献 → 写一节 → 导出」"],
            ["C  对外试点", "第二个实验室", "对方自己导入文献并写出一节，附一页运维说明"],
        ],
        0.48,
        1.4,
        12.35,
        3.7,
        [3.15, 3.4, 5.8],
        13,
    )
    rect(slide, 0.48, 5.3, 12.35, 1.55, SAND)
    txt(slide, "现在卡在 A 的后半句", 0.72, 5.48, 11.9, 0.32, 14, True, RUST)
    txt(
        slide,
        "工程上系统已经能走完，三篇学长稿把「AI 能写到哪」标定了。缺的是指定的人用真题自己走完。第一篇作者自走的稿出来之前，不用「等反馈」去开 LaTeX、基金标书、多实验室网站。",
        0.72,
        5.88,
        11.9,
        0.8,
        14,
        False,
        INK,
    )


def foundations(prs) -> None:
    slide = page(prs, 9, "基础", "四件实验室原来就在做的事", "后面四页各讲一件。")
    table(
        slide,
        [
            ["事项", "以前", "现在"],
            ["查重", "按次付费，问题段自己找", "组内预检 + 降重定位，再决定送不送终检"],
            ["文献", "PDF 在个人电脑，人走库散", "组内库增量入库，写作时按章节检索"],
            ["图表", "Origin、Jade、PPT 来回切", "同一工作台出图、出表，进文稿"],
            ["上手", "按钮多，不知道先点哪", "跟 Agent 说下一步；关键步骤仍等人"],
        ],
        0.48,
        1.4,
        12.35,
        4.55,
        [1.8, 4.7, 5.85],
        14,
    )
    txt(slide, "这四件都是课题组原来就要花钱或花时间的。系统把它们收进同一条写稿路径。", 0.48, 6.15, 12.35, 0.5, 14, False, MUTED)


def plagiarism(prs) -> None:
    slide = page(prs, 10, "查重", "从反复送检，变成组内先看", "不能代替知网终检。预检分数不能当成学校分数。")
    rect(slide, 0.48, 1.4, 6.05, 3.55, WHITE, LINE)
    txt(slide, "上学期的真实开支", 0.72, 1.58, 5.55, 0.38, 15, True, RUST)
    lines(slide, ["初稿测、改完测、临近提交再测", "单次单价现在并不夸张", "贵在往返次数", "还没定位问题段就送正式库"], 0.72, 2.15, 5.55, 2.5, 15)
    rect(slide, 6.73, 1.4, 6.1, 3.55, NAVY)
    txt(slide, "系统里现在做的", 6.97, 1.58, 5.6, 0.38, 15, True, RUST)
    lines(
        slide,
        ["相似性预检，标出重复段落", "降重后并排看，再决定采纳", "和质量审查放在同一个入口", "写完本节就能去"],
        6.97,
        2.15,
        5.6,
        2.5,
        15,
        WHITE,
    )
    rect(slide, 0.48, 5.15, 12.35, 1.7, SAND)
    txt(slide, "边界", 0.72, 5.32, 11.9, 0.32, 14, True, NAVY)
    txt(
        slide,
        "预检用来减少正式库的往返。学校或期刊指定的知网 / 维普终检仍然要做。若被问能不能代替知网：不能代替终检，能少做几次无准备的付费检测。",
        0.72,
        5.72,
        11.9,
        0.9,
        14,
        False,
        INK,
    )


def library(prs) -> None:
    slide = page(prs, 11, "知识库", "文献从个人盘变成组里的增量库", "稿好不好，下一步看库新不新。每方向千篇是工作目标，不是已经自动达成的数字。")
    cells = [
        ("留在组里", "私域 PDF 进组内索引。换届、换电脑，库不断。检索是关键词加语义，两路合并。"),
        ("按章检索", "按论文结构切块。写引言、方法、结果时，检索偏向不同。"),
        ("只处理新增", "新论文不必整库重做。题录可走 RIS / BibTeX，开放获取全文可入库。"),
        ("大约 900 篇", "历史上曾经到这个量级。之后靠各方向继续投喂，库才会新。"),
    ]
    for i, (h, d) in enumerate(cells):
        col, row = i % 2, i // 2
        x, y = 0.48 + col * 6.35, 1.4 + row * 2.55
        rect(slide, x, y, 6.15, 2.38, WHITE, LINE)
        rect(slide, x, y, 0.12, 2.38, RUST if i == 0 else NAVY)
        txt(slide, h, x + 0.4, y + 0.28, 5.5, 0.45, 18, True, NAVY)
        txt(slide, d, x + 0.4, y + 0.9, 5.5, 1.15, 14, False, INK)


def charts(prs) -> None:
    slide = page(prs, 12, "图表", "少切几次软件", "覆盖本实验室常用科研图。特殊期刊精排仍要作者收尾。")
    txt(slide, "以前：表格 → Origin / Jade → PPT → 分子网站 → 贴回 Word。", 0.48, 1.38, 12.35, 0.38, 14, False, MUTED)
    kinds = [
        ("数据图", "柱状、折线、散点、热力、面积"),
        ("表征与计算", "XRD、DFT"),
        ("三线表", "统计结果直接进文稿"),
        ("示意图", "流程、机理草图、分子结构"),
    ]
    for i, (h, d) in enumerate(kinds):
        x = 0.48 + i * 3.15
        rect(slide, x, 1.9, 3.02, 2.55, WHITE, LINE)
        rect(slide, x, 1.9, 3.02, 0.12, RUST if i % 2 else NAVY)
        txt(slide, h, x + 0.15, 2.25, 2.72, 0.55, 17, True, NAVY, PP_ALIGN.CENTER)
        txt(slide, d, x + 0.15, 2.95, 2.72, 1.1, 13, False, MUTED, PP_ALIGN.CENTER)
    rect(slide, 0.48, 4.7, 12.35, 2.15, NAVY)
    txt(
        slide,
        "图在同一注册表里增加。写作和 Agent 可以调用，经过布局和质检后再进文稿，导出带题注清单。龙智、陈韶光、易程三篇原稿里都已经进过图。",
        0.75,
        5.15,
        11.85,
        1.25,
        16,
        False,
        WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def agent(prs) -> None:
    slide = page(prs, 13, "上手", "说下一步，关键步等人", "Agent 省找功能的时间，省不了科研判断。")
    rect(slide, 0.48, 1.4, 6.05, 3.5, WHITE, LINE)
    txt(slide, "以前", 0.72, 1.58, 5.55, 0.35, 15, True, MUTED)
    lines(slide, ["功能齐，路径像填题、点生成、导出", "作者参与少", "新人不知道该点哪个"], 0.72, 2.15, 5.55, 2.4, 15)
    rect(slide, 6.73, 1.4, 6.1, 3.5, NAVY)
    txt(slide, "现在", 6.97, 1.58, 5.6, 0.35, 15, True, RUST)
    lines(
        slide,
        ["跟 Agent 说下一步", "系统去调文献、写节、作图、检查", "每一轮留下一个看得见的结果", "没有大纲会先问，不会直接往下写"],
        6.97,
        2.15,
        5.6,
        2.4,
        15,
        WHITE,
    )
    rect(slide, 0.48, 5.1, 12.35, 1.75, SAND)
    txt(slide, "仍然等人点头的地方", 0.72, 5.28, 11.9, 0.32, 14, True, NAVY)
    txt(
        slide,
        "大纲和蓝图必须批准。导师确认过的方向是人机协作：人先写要点，系统再扩写。三篇学长稿是开发者带着系统和材料代跑出来的——流程能走通，判断仍在人。",
        0.72,
        5.68,
        11.9,
        0.95,
        14,
        False,
        INK,
    )


def path(prs) -> None:
    slide = page(prs, 14, "本阶段", "一条能走完的初稿路径", "这一页证明框架有了。下一页用三篇真人稿证明上限看见了。")
    steps = [
        "1  方向或新建项目",
        "2  导入文献，确认配置",
        "3  大纲批准",
        "4  按节起草并质检",
        "5  实验数据进入结果",
        "6  图和表进文稿",
        "7  审查、查重、降重",
        "8  导出 Word / PDF",
    ]
    for i, step in enumerate(steps):
        col, row = i % 4, i // 4
        x, y = 0.48 + col * 3.15, 1.45 + row * 2.35
        fill = NAVY if i == 7 else WHITE
        fg = WHITE if i == 7 else NAVY
        rect(slide, x, y, 3.02, 2.15, fill, None if i == 7 else LINE)
        txt(slide, step, x + 0.15, y + 0.7, 2.72, 0.75, 15, True, fg, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def errors(prs) -> None:
    slide = page(prs, 15, "质量", "先拦住三类最费时间的错", "不保证句子漂亮。贡献点立不立得住，要作者判。")
    items = [
        ("乱引", "引用编号只能来自检索到的文献，越界会去掉并告警。"),
        ("编造数字", "精确数字要对上实验数据或文献，对不上就告警。"),
        ("结构串章", "结果和讨论分开写。章节之间做一致性检查。写完可以进审查清单。"),
    ]
    for i, (h, d) in enumerate(items):
        y = 1.4 + i * 1.7
        rect(slide, 0.48, y, 12.35, 1.55, WHITE, LINE)
        rect(slide, 0.48, y, 2.55, 1.55, NAVY)
        txt(slide, h, 0.6, y + 0.5, 2.3, 0.55, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, d, 3.3, y + 0.4, 9.2, 0.75, 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)


def case(prs) -> None:
    slide = page(prs, 16, "案例", "三篇学长初稿，上限在哪", "都是代跑。当场打开一份 Word。终稿仍是作者的。")
    table(
        slide,
        [
            ["学长", "写什么", "已经能到", "到不了（上限）"],
            ["龙智", "茶树 CSS0022168 负调控茶氨酸", "结构、数字、图能进第一稿", "新梢发育有没有直接表型；投稿英文"],
            ["陈韶光", "红麻催化热解，生物炭 / HZSM-5", "问题→结果→机理几天内能铺开", "单次实验、峰面积分数、缺对照，机理会被写圆"],
            ["易程", "KSC3 生物炭降茶树氟（多品种）", "降氟、品质、微生物可给导师看", "同一题两种贡献句都写得出，必须人锁"],
        ],
        0.4,
        1.38,
        12.5,
        3.95,
        [1.4, 3.5, 3.7, 3.9],
        12,
    )
    rect(slide, 0.48, 5.5, 12.35, 1.35, SAND)
    txt(
        slide,
        "第一版远快于从零手写；组内数据和引用能进稿；贡献点立不立得住，系统替代不了。下一期不是再代写一篇，是作者自己打开系统走一遍。原稿在 F:\\论文。",
        0.72,
        5.7,
        11.9,
        0.95,
        14,
        False,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def artifacts(prs) -> None:
    slide = page(prs, 17, "证据", "会上带三样实物", "没有第 1 样，后两样撑不住。")
    items = [
        ("01", "三份初稿", "龙智带图稿、陈韶光 CEJ 中文稿、易程降氟全文。路径：F:\\论文\\龙智、陈韶光、易程。"),
        ("02", "一页过程", "用了哪些步骤、卡在哪、人工改了什么、接口大约花了多少。写清代跑 vs 作者自己点。"),
        ("03", "一页对照", "ChatGPT、Paperpal、本系统各做哪一段。本系统不做什么，写在同一页。"),
    ]
    for i, (n, h, d) in enumerate(items):
        y = 1.4 + i * 1.7
        rect(slide, 0.48, y, 12.35, 1.55, WHITE, LINE)
        rect(slide, 0.48, y, 1.55, 1.55, NAVY)
        txt(slide, n, 0.48, y + 0.5, 1.55, 0.55, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, h, 2.25, y + 0.22, 10.2, 0.4, 18, True, NAVY)
        txt(slide, d, 2.25, y + 0.72, 10.2, 0.6, 14, False, MUTED)


def boundary(prs) -> None:
    slide = page(prs, 18, "边界", "做到的，和先写明的", "边界自己先说，比被问到再解释更像做过取舍。")
    rect(slide, 0.48, 1.4, 6.05, 5.45, WHITE, LINE)
    rect(slide, 0.48, 1.4, 6.05, 0.55, NAVY)
    txt(slide, "这期做到", 0.72, 1.5, 5.55, 0.38, 16, True, WHITE)
    lines(
        slide,
        [
            "私域检索、分节写作、引用约束",
            "审查和查重预检、常用科研图",
            "Word / PDF，关键步骤等人批准",
            "实验室可以部署试用",
            "三篇真实题目的可审校初稿（代跑）",
            "上限可以指给老师看",
        ],
        0.72,
        2.2,
        5.55,
        4.3,
        15,
    )
    rect(slide, 6.73, 1.4, 6.1, 5.45, WHITE, LINE)
    rect(slide, 6.73, 1.4, 6.1, 0.55, RUST)
    txt(slide, "不放进这期承诺", 6.97, 1.5, 5.6, 0.38, 16, True, WHITE)
    lines(
        slide,
        [
            "完整产品；一键全文；无人值守可投稿",
            "期刊级精排、稳定的 LaTeX",
            "二十人同时写、全年无运维",
            "替代知网终检",
            "生信平台或基金标书系统",
            "再做一套站点，或把管线推倒重写",
        ],
        6.97,
        2.2,
        5.6,
        4.3,
        15,
    )


def next_phase(prs) -> None:
    slide = page(prs, 19, "下一阶段", "4～8 周：先打穿真题", "工程用一句带过。时间留给真题怎么组织。")
    rect(slide, 0.48, 1.4, 7.55, 5.45, WHITE, LINE)
    txt(slide, "主轴：试用", 0.72, 1.58, 7.1, 0.38, 16, True, RUST)
    lines(
        slide,
        [
            "定性：驻组写作系统，验收可审校初稿",
            "选 1～2 个方向，各 1 篇真题",
            "作者自己点，开发者陪跑",
            "空项目 → 文献 → 大纲批准 → 写节",
            "图 / 表 → 质量中心 → 导出",
            "回收：质量、速度、卡点、敢不敢投稿前自查",
            "只按清单改小处，事先说好上限",
        ],
        0.72,
        2.15,
        7.05,
        4.4,
        14,
    )
    rect(slide, 8.23, 1.4, 4.6, 5.45, NAVY)
    txt(slide, "同时做的工程", 8.48, 1.58, 4.15, 0.4, 15, True, RUST)
    lines(
        slide,
        ["页面内容消毒", "中断时后台也能停", "作图和上传收紧", "建项和自动保存不丢", "工作台编排下沉", "LaTeX 未拍板不开"],
        8.48,
        2.2,
        4.15,
        4.2,
        14,
        WHITE,
        10,
    )


def difficulty(prs) -> None:
    slide = page(prs, 20, "难点", "主因不是学不会，是没有截止日期", "学不会和不急着写要分开说。")
    rect(slide, 0.48, 1.4, 8.05, 5.45, WHITE, LINE)
    txt(slide, "现实层 · 主因", 0.72, 1.58, 7.55, 0.32, 13, True, RUST)
    txt(slide, "产品定性没锁死；组里现在不急着写。", 0.72, 2.0, 7.55, 0.45, 18, True, NAVY)
    lines(
        slide,
        [
            "没有「不写不行」的节点，任何写作工具都是选修课。",
            "三篇学长稿是代跑，填不上这条空档。",
            "下一阶段先绑到已有任务：",
            "·  开题或中期必须交的综述一章",
            "·  基金本子里的已有工作，改写成引言",
            "·  一组数据，两周内只走工作台，出结果节骨架",
        ],
        0.72,
        2.6,
        7.55,
        3.9,
        14,
    )
    rect(slide, 8.73, 1.4, 4.1, 5.45, NAVY)
    txt(slide, "技术层 · 能扛", 8.97, 1.58, 3.65, 0.35, 13, True, RUST)
    lines(
        slide,
        ["安全、中断、保存不丢", "人多时 4 核 8G 会慢", "库不再入库就会旧", "开发仍是一人", "瓶颈在试用组织", "不先招第二名全职开发"],
        8.97,
        2.15,
        3.65,
        4.3,
        14,
        WHITE,
        8,
    )


def support(prs) -> None:
    slide = page(prs, 21, "支持", "请老师给使用闭环", "缺的不是开发费。缺的是定性、人、题、看稿日。")
    asks = [
        ("1", "一句话定性", "驻组写作系统；作者主笔；验收可审校初稿。不验收完整产品，不验收一键全文。"),
        ("2", "一个真实写稿任务", "作者、题目、交稿日、老师看稿日，填在下一页。绑到开题 / 中期 / 两周结果节，不要另起炉灶。"),
        ("3", "再点 1～2 个人走完一节", "2～4 周完成「文献 → 写一节 → 导出」。我陪跑，作者自己点。收回一张卡点清单。"),
    ]
    for i, (n, h, d) in enumerate(asks):
        y = 1.38 + i * 1.28
        rect(slide, 0.48, y, 12.35, 1.16, WHITE, LINE)
        rect(slide, 0.48, y, 0.85, 1.16, RUST if i == 1 else NAVY)
        txt(slide, n, 0.48, y + 0.3, 0.85, 0.55, 22, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, h, 1.55, y + 0.12, 10.9, 0.35, 16, True, NAVY)
        txt(slide, d, 1.55, y + 0.52, 10.9, 0.5, 13, False, MUTED)
    rect(slide, 0.48, 5.3, 12.35, 1.55, SAND)
    txt(
        slide,
        "维持：DeepSeek / 智谱和 4 核 8G 按现有规模。文献继续入库。不请大额开发费，不请第二名全职开发。LaTeX、基金本子、第二个实验室这期不做。不要用「组员去学」代替交稿日。",
        0.72,
        5.5,
        11.9,
        1.15,
        14,
        False,
        INK,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def close(prs) -> None:
    slide = page(prs, 22, "拍板", "请当场填四格", "老师不填第 1、第 2 格，后面两格下次再说。")
    rows = [
        ("1", "定性", "□ 驻组、可审校初稿          □ 其他 ________________"),
        ("2", "第一篇作者自走", "作者 ______    题目 ______    交稿日 ______    老师看稿日 ______"),
        ("3", "再走完一节的人", "______________            ______________"),
        ("4", "模型 / LaTeX", "□ 维持现有     □ 按 5 人试用上调     □ LaTeX 这期不做"),
    ]
    for i, (n, h, d) in enumerate(rows):
        y = 1.38 + i * 1.05
        rect(slide, 0.48, y, 12.35, 0.95, WHITE, LINE)
        rect(slide, 0.48, y, 0.7, 0.95, NAVY)
        txt(slide, n, 0.48, y + 0.22, 0.7, 0.5, 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        txt(slide, h, 1.4, y + 0.08, 11.1, 0.3, 13, True, RUST)
        txt(slide, d, 1.4, y + 0.42, 11.1, 0.4, 14, False, INK)
    txt(slide, "请用一篇作者自己走的真题，给驻组迭代一个闭环。其他实验室是这条路走通之后的事。", 0.48, 5.65, 12.35, 0.5, 14, False, MUTED)


if __name__ == "__main__":
    build()
