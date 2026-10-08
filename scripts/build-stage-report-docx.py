"""Briefing Word for advisors — narrative report, not a slide-script dump.

Usage: python scripts/build-stage-report-docx.py
Output: docs/reports/禾书耕文-阶段汇报.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports" / "禾书耕文-阶段汇报.docx"

NAVY = RGBColor(0x14, 0x2A, 0x45)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)


def font(run, name="宋体", size=12, bold=False, color=INK, east=None):
    east = east or name
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:cs"), name)


def p(doc, text, *, size=12, bold=False, center=False, first=True, space_before=0, space_after=8, color=INK, name="宋体"):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.first_line_indent = Cm(0.74) if first and not center else Cm(0)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    font(run, name=name, east="黑体" if name == "黑体" else name, size=size, bold=bold, color=color)
    return para


def h(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run(text)
    font(run, name="黑体", east="黑体", size=sizes[level], bold=True, color=NAVY)


def bullets(doc, items: list[str]):
    for item in items:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.74)
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run("·　" + item)
        font(run, size=12)


def shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "888888")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_table(doc, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            run = para.add_run(text)
            header = r == 0
            font(run, name="宋体", east="宋体", size=10.5, bold=header or c == 0, color=RGBColor(0xFF, 0xFF, 0xFF) if header else INK)
            shade(cell, "142A45" if header else ("F4F1EA" if r % 2 else "FFFFFF"))
            borders(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p(doc, "禾书耕文", size=22, bold=True, center=True, first=False, name="黑体", color=NAVY, space_after=4)
    p(doc, "实验室科研写作系统　阶段汇报", size=16, bold=True, center=True, first=False, name="黑体", color=NAVY, space_after=14)
    p(doc, "（供课题组导师 / 智慧农业创新中心审阅）", size=12, center=True, first=False, color=MUTED, space_after=18)

    meta = [
        "汇报日期：2026年9月30日（2026年10月4日按融合口径修订）",
        "项目：农业科研 AI 辅助写作（私域文献 + 人机协作，作者主笔）",
        "汇报人：________________",
        "稿件样本：F:\\论文\\龙智、陈韶光、易程",
    ]
    for line in meta:
        p(doc, line, first=False, size=12, space_after=4)

    h(doc, "一、请示事项")
    p(
        doc,
        "本期工程框架已经能够走出一篇可审校的论文初稿。下一期不请求新功能清单，也不请求第二名全职开发。请老师决定：是否将本系统定性为课题组驻组写作工具，并指定真实写稿任务，形成使用闭环。",
    )
    p(doc, "具体请予明确的事项如下：")
    bullets(
        doc,
        [
            "定性：驻组写作系统，作者主笔，验收口径为「可审校初稿」，不验收完整产品，不验收一键生成全文。",
            "任务：指定一名作者、一道真题、一个交稿日、老师看稿日期；尽量绑到已有节点（开题、中期或两周内的结果节骨架）。",
            "复制：再指定一至两人，在二至四周内独立完成「文献 → 写一节 → 导出」。开发者陪跑，作者自己操作。",
            "运转：DeepSeek、智谱接口与现有四核八吉服务器按小规模维持；文献继续入库。不请大额开发费。",
        ],
    )
    p(doc, "文末附填空表，便于会上直接填写。")

    h(doc, "二、阶段判断")
    p(
        doc,
        "系统已从「功能较全的写作工具箱」走到「能走完一篇论文初稿的实验室系统」。路径为：方向规划 → 建立项目 → 工作台分节写作 → 引用、审查、查重预检、配图 → 导出 Word / PDF。关键步骤需作者批准后才进入下一步。写作管道为起草（DeepSeek）→ 核对（智谱，失败时有备用）→ 修改。引用越界会去掉并告警。",
    )
    p(
        doc,
        "证据有两层。系统层：查重预检、组内文献库、工作台配图、按自然语言推进的 Agent，已经串在同一条写稿路径上。稿件层：已为龙智、陈韶光、易程三位学长代跑真实题目初稿，用来标明现阶段人工智能初稿的上限，也使开发者把论文从材料到成稿的流程走通。",
    )
    p(
        doc,
        "尚未证明的是：组员自己打开系统、带着真题走完。三篇学长稿是代跑样本，不是规模化使用证明。仓库约四百余次提交。近期收口是每轮只交出一个可见结果，蓝图与大纲须人批准。",
    )

    h(doc, "三、与现有工具的关系")
    p(
        doc,
        "公网工具大致四类：润色与投稿检查（如 Paperpal）、海量文献助手（如 SciSpace）、综述 Agent（如 Elicit）、通用对话（ChatGPT 及国产对话）。国内还有功能更多、运营更强的科研平台。功能个数和品牌不是本项目的比较对象，以免被理解成「如何胜过解螺旋」。",
    )
    p(doc, "公开标价仅作量级对照（约 2026 年，非合同价）：")
    bullets(
        doc,
        [
            "Paperpal：Prime 约 144 美元/人/年，Pro 约 348 美元/人/年。",
            "SciSpace Premium：年付约合 144 美元/年；文献任务重时更高。",
            "Elicit Pro：约 49 美元/人/月（年付），偏系统综述。",
            "五人组若每人订一份润色或文献助手，一年大约几千到一万出头人民币。文献和实验数据多在对方云上，或根本进不了组内稿。",
            "查重：知网个人约 1.5 元/千字，AIGC 检测另计约 2 元/千字。硕士稿约三万字，一轮查重加 AIGC 大约一百出头；改一版再测，一学期可到数百。学校或期刊指定的终检不能省，能省的是尚未改完就反复送正式库。",
        ],
    )
    p(
        doc,
        "本系统要维持的开支，与五人组一年公网订阅同一量级：服务器、模型接口、陪跑。比较对象是课题组已经在花的钱，不是上市公司估值。下一期也不是再做一个完整产品，而是把组内路径绑到真实写稿任务上。",
    )
    p(doc, "我们占住的四条差异是：")
    bullets(
        doc,
        [
            "组内 PDF、实验表和稿留在实验室服务器。",
            "写作落到当前章节，引用不得超出本次检索到的文献。",
            "结果章须接得上实验数据，数字不能凭对话编造。",
            "农科写法写进流程：IMRaD、结果与讨论分开，过度绝对的措辞会拦截。",
        ],
    )
    p(
        doc,
        "公网订阅仍可用于英文润色和海量检索。禾书耕文补的是它们进不了组内数据和组内文库的那一段。不是全面替代。",
    )

    h(doc, "四、产品定位与核心优势")
    p(
        doc,
        "面向农业课题组的私有部署写作流程。作者主笔。系统负责组内检索、按章起草、配图、审查、查重预检，以及导出可交导师审阅的 Word / PDF。学生按阶段推进；大纲、蓝图和每一轮结果，须作者点头后才进入下一步。",
    )
    p(
        doc,
        "核心优势需要讲清楚：本汇报人不能保证做出一个完整的、别人拿去就会用的产品。能够保证的是人一直在课题组里，按真实写稿需求改，做成最贴本组习惯的系统。这是驻组工程，不是外包交钥匙，也不是用功能清单去打公网品牌。听到「一键出全文」，应回到本段：交出的是可审校初稿。三年到五年的国内推广是后续叙事。本期只要求：先成为本实验室比直接打开对话框更值得使用的工具。",
    )

    h(doc, "五、分期验收")
    p(doc, "本期经费不按终局验收。分三期，每期有证据，过不了则停止追加开发投入。")
    add_table(
        doc,
        [
            ["期", "本期换到什么", "怎样算过"],
            ["A　本实验室打穿", "算力与陪跑", "系统能走完；三篇代跑可看上限；另有一篇作者自己操作的初稿，导师看过"],
            ["B　组内能复制", "按问题清单小改，有上限", "至少三人独立完成「文献 → 写一节 → 导出」"],
            ["C　对外试点", "第二个实验室", "对方自行导入文献并写出一节，附一页运维说明"],
        ],
    )
    p(
        doc,
        "工程上已接近 A 的「系统能走完」。三篇学长稿标定了「人工智能能写到哪」。缺的是 A 的「人真正自己走完」。第二个实验室须等本实验室有人能独立走完再议。细功能等实际使用再改，这一判断仍然成立。在作者自己走出第一篇真稿之前，不以「等待反馈」为由开启 LaTeX、基金标书或多实验室网站，否则「等使用」会变成无限延期。",
    )

    h(doc, "六、已经收进同一条路径的四项基础")
    p(doc, "下列四项本是课题组原来就要花钱或花时间的事，现已收进同一写稿路径。")
    add_table(
        doc,
        [
            ["事项", "以前", "现在"],
            ["查重", "按次付费，问题段自己找", "组内预检与降重定位，再决定是否送终检"],
            ["文献", "PDF 在个人电脑，人走库散", "组内库增量入库，写作时按章节检索"],
            ["图表", "Origin、Jade、PPT 等来回切换", "同一工作台出图出表，进入文稿"],
            ["上手", "按钮多，不知先点哪", "用一句话说下一步；关键步骤仍须等人"],
        ],
    )

    h(doc, "（一）查重：组内先看，终检仍做", 2)
    p(
        doc,
        "以往初稿、改完、临近提交往往各测一次。单次单价并不夸张，贵在往返，以及尚未定位问题段就送正式库。系统现有相似性预检、降重对照、与审查同处的质量中心。预检用于减少无准备的付费检测，不能代替学校或期刊指定的知网、维普终检，预检分数也不能当成学校分数。",
    )

    h(doc, "（二）文献：组内增量库", 2)
    p(
        doc,
        "私域 PDF 进入组内索引，关键词与语义两路合并。按论文结构切块，写引言、方法、结果时检索偏向不同。新增文献只处理新增部分。外部题录可走 RIS / BibTeX，开放获取全文可入库。历史上索引约九百篇量级，此后须各方向继续投喂。中心提出的「每方向约千篇」是文献工作目标，不是系统已经自动达成的数字。稿件质量下一步主要取决于库是否更新、方向是否齐全、实验表是否完整，而不是再换一个更大的模型。",
    )

    h(doc, "（三）图表：少切换软件", 2)
    p(
        doc,
        "同一注册表可出常用数据图、XRD 与 DFT 等表征计算图、三线表，以及流程、机理草图、分子结构等示意图。写作和 Agent 可调用，经布局和质检后进入文稿，导出带题注清单。三位学长稿中均已进图。覆盖本实验室常用科研图；特殊期刊精细排版仍须作者收尾。",
    )

    h(doc, "（四）上手：说下一步，关键步等人", 2)
    p(
        doc,
        "导师此前已确认人机协作方向：人先写要点和大纲，系统再扩写。现在默认入口是 Agent：学生说明下一步，系统去检索、写节、作图、检查，每轮留下可见结果；无大纲则先询问。学习成本仍在，但位置变了：不再先学十几个按钮，而是会改系统交出的那一节，并对数字和引用负责。Agent 节省的是找功能的时间，节省不了科研判断。三篇学长稿由开发者代跑，也说明流程能走通，判断仍在人。",
    )

    h(doc, "七、本阶段已结束的初稿路径与质量闸门")
    p(doc, "作者可按下列检查点推进：新建或选择方向与项目；导入文献并确认配置；批准大纲；按节起草、核对、修改；实验数据进入结果；图表进文稿；质量中心审查、查重预检与降重；导出前就绪检查后输出 Word / PDF（含双语摘要与图表题注）。实验室已有登录与管理、用量记录，并在约四核八吉的机器上做过小规模试跑。")
    p(doc, "与直接把题目交给对话模型相比，流程中硬性拦住三类导师改稿最耗时的错误：引用越界、编造或对不上的精确数字、结果与讨论串章及过度绝对措辞。不保证句子漂亮。贡献点是否立得住，系统替代不了，须作者判断。")

    h(doc, "八、三篇学长初稿所标明的上限")
    p(
        doc,
        "三篇均为开发者带着系统和材料代跑，原稿在 F:\\论文。目的是让组里看见现阶段人工智能初稿能到哪里、不能到哪里，并让开发者走通全流程。不是组员已经能够独立使用。所指周期是可审校初稿，终稿、投稿和返修仍由作者完成。",
    )
    add_table(
        doc,
        [
            ["学长", "题目与稿型", "已经能够做到", "到不了（上限）"],
            [
                "龙智",
                "茶树 CSS0022168 负调控茶氨酸（英文研究论文，带图）",
                "结构、数字和图可进入第一稿",
                "新梢发育有无直接表型、投稿级英文；作者不核对即会错",
            ],
            [
                "陈韶光",
                "红麻催化热解，生物炭与 HZSM-5 双轨升级（化工环境类）",
                "问题、结果、机理的骨架可在较短时间内铺开",
                "单次实验、峰面积分数、缺对照等，系统容易把机理写圆",
            ],
            [
                "易程",
                "KSC3 生物炭降低茶树氟积累（多品种盆栽）",
                "降氟、品质、微生物可写成提交导师的初稿",
                "同一题目可写出两种贡献句（如茶氨酸大幅上升与幼叶基本稳定），必须由作者锁定",
            ],
        ],
    )
    p(
        doc,
        "三篇合在一起的结论是：第一版叙述远快于从零手写；组内数据和引用能够进稿；贡献点是否成立，系统替代不了。下一期需要的不是再代写一篇，而是作者自己打开系统走一遍。会上可打开其中一份原稿。建议同时准备一页过程记录（步骤、卡点、人工修改）和一页与 ChatGPT、Paperpal 的分工对照。代码提交次数和人天估价只能说明沉没成本，不能单独作为追加经费的理由。",
    )

    h(doc, "九、本期边界")
    p(doc, "已经做到：私域检索、分节写作、引用约束、审查和查重预检、常用科研图、Word / PDF；关键步骤须人批准；实验室可以部署试用；三篇真实题目的可审校初稿（代跑）可供对照上限。")
    p(doc, "不列入本期承诺：完整产品；无人值守、可直接投稿的全文；期刊级精细排版和稳定的 LaTeX；二十人同时写作、全年无运维；替代知网终检；做成生信平台或基金标书系统；再做一套站点或推倒重写写作管线。")

    h(doc, "十、推进难点")
    p(
        doc,
        "学习成本是存在的。新人须明白：系统交出的是草稿，数字和引用须自己核对，质量中心是预检。但这不是主因。",
    )
    p(
        doc,
        "技术层面可以继续做：页面安全、写作中断、作图与上传进程、保存不丢、工作台减负。多人同时写时现有机器容易变慢，此前已做限流和超时修复，扩大试用需升配或错峰。库若不再入库，检索会旧。开发仍是一人。瓶颈在试用组织，不在再招一名全职开发。",
    )
    p(
        doc,
        "现实层面才是主因。其一，产品定性一直没有锁死，工具箱、写作流程、Agent、查重替代混在一起，验收对不齐，驻组迭代容易变成每问一个功能就加一个。其二，组里目前并不急着写论文，使用积极性不高；没有「不写不行」的截止日期，任何写作工具都是选修课。三篇代跑填不上这条空档。",
    )
    p(doc, "因此下一阶段（建议四至八周）不先做厚培训手册，而把系统绑到已经存在的任务上，例如开题或中期综述中的一章、基金本子「已有工作」改写成引言、老师指定一组数据两周内只走工作台交出结果节骨架。没有指定题目和交稿日，按使用迭代就是空转。")

    h(doc, "十一、请予支持的事项")
    p(doc, "现在缺少的不是第二名开发，也不是大额开发费。缺少的是使用闭环。")
    h(doc, "（一）主请求", 2)
    add_table(
        doc,
        [
            ["序号", "请予决定", "具体程度", "用以换取"],
            ["1", "一句话定性", "驻组、作者主笔、可审校初稿；不验收完整产品和一键全文", "后续改功能有尺子"],
            ["2", "一个真实写稿任务", "作者、题目、交稿日、老师看稿日；绑到已有节点", "截止日前必须有人打开系统，按卡点修改"],
            ["3", "再点一至两人走完一节", "二至四周「文献 → 写一节 → 导出」；作者自己点，开发者陪跑", "证明并非只有开发者会用，并回收卡点清单"],
        ],
    )
    h(doc, "（二）维持运转（不是新项目）", 2)
    add_table(
        doc,
        [
            ["项目", "请求", "不请求"],
            ["模型", "DeepSeek、智谱按现有规模稳住；五人同时试用再议", "不按全国实验室加额度"],
            ["机器", "四核八吉够小规模；人多再升配；文献盘和备份指定到人", "不先购置大型机器"],
            ["文献", "各方向继续向组内库投放 PDF，并写明使用范围", "不要求系统自动凑齐每方向一千篇"],
            ["学术口径", "人工智能稿必须人工审校；数据与引用由作者负责", "不另做一套伦理系统"],
            ["经费", "接口费、陪跑、按清单改小处", "不请大额开发费，不请第二名全职开发"],
        ],
    )
    h(doc, "（三）请明确本期不做", 2)
    bullets(
        doc,
        [
            "LaTeX：默认这四至八周不做；若要做，另开一次，先定期刊模板。",
            "基金标书、生信平台、第二个实验室：待本实验室有人自己走完再议。",
            "不以「让组员去学」代替指定交稿日。学不会不是主因，不急着写才是。",
        ],
    )

    h(doc, "十二、请会上填写")
    p(doc, "请老师优先填写第 1、第 2 项。两项不定，其余可下次再议。")
    add_table(
        doc,
        [
            ["序号", "填写"],
            ["1　定性", "□ 驻组、可审校初稿　　□ 其他：____________________"],
            ["2　第一篇作者自走", "作者________　题目________　交稿日________　老师看稿日________"],
            ["3　再走完一节的人", "______________　　______________"],
            ["4　模型与 LaTeX", "□ 维持现有　　□ 按五人试用上调　　□ LaTeX 本期不做"],
        ],
    )
    p(
        doc,
        "以上汇报请予审阅。框架能够出初稿，三篇原稿把上限摊开了。请用一篇作者自己完成的真题，给驻组迭代一个闭环。其他实验室是这条路走通之后的事情。",
        first=True,
    )

    h(doc, "附录一　三个常问问题")
    p(doc, "与 ChatGPT 有何不同？文献和数据在组里，引用不能越出文献池，结果章须接实验数据，关键步骤作者点头。流畅段落系统写得出，是否立得住仍要作者判断。", first=True)
    p(doc, "为何还要投入？主请求不是钱，是指定人和交稿日。钱只维持接口和机器并用于陪跑，换取作者自己点出的初稿和一张卡点清单。", first=True)
    p(doc, "能否推广到别的实验室？可作为后续目标。须先有本实验室第二人能够不靠手把手完成「文献 → 写一节 → 导出」。", first=True)

    h(doc, "附录二　投入口径（内部对照，非审计价）")
    p(
        doc,
        "2026年6月1日内部估算约一百八十至二百二十人天，按外包中位人天对外常报二十五万至四十万（见项目估值说明）。此后七至九月的 Agent、写作与图表质量系统、质量中心，以及三篇题目代跑，未计入该表，实际投入更高。该组数字只说明沉没成本已经不小。追加投入的理由只能是：下一期能够换到作者自己走完的初稿和问题清单。",
    )

    h(doc, "附录三　前期工作一览")
    add_table(
        doc,
        [
            ["阶段", "时间（约）", "主要工作"],
            ["起步与初版", "2026年5月", "认证、工作台、写作管道、图表与 XRD、文献对话、农学规范；开始部署"],
            ["工程化", "2026年5月末至6月初", "数据库、检索索引、管理端、安全与限流、质量查重审查"],
            ["产品纠偏", "2026年6月", "由一点生成转为作者先写要点和大纲；导师确认人机协作"],
            ["文献与方向", "2026年6月至7月", "期刊信息、外部检索、题录入库；研究方向模块"],
            ["生命周期", "2026年7月", "安全收拢、项目阶段、Agent、引用硬检、审查轮次、导出就绪"],
            ["Agent 产品化", "2026年7月末至9月", "写作与图表质量、附件与数据、质量中心、增量检索、每步等人"],
            ["真实题目代跑", "2026年6月至9月", "龙智、陈韶光、易程三篇学长初稿"],
        ],
    )
    p(
        doc,
        "试用中已处理过的非正式反馈包括：多人同时写变慢、路径像填题生成、文献需要期刊信息、写完乱跳、改引用反复、上传表格解析失败。已对应做过限流、人机协作、外部文献导入、过程可见和附件解析。中心提出的「各方向入库并用本系统写一篇」在系统上走得通；论文与文献配额的完成情况不在本仓库统计范围内。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
