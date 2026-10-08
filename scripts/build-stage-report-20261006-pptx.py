"""2026-10-06 stage report deck from the 14-page outline.

Usage: python scripts/build-stage-report-20261006-pptx.py
Output: docs/reports/阶段汇报-2026-10-06.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports" / "阶段汇报-2026-10-06.pptx"

FONT = "Microsoft YaHei"
NAVY = RGBColor(0x14, 0x2A, 0x45)
NAVY_SOFT = RGBColor(0x1E, 0x3D, 0x5C)
RUST = RGBColor(0xC0, 0x56, 0x2A)
RUST_SOFT = RGBColor(0xF8, 0xEB, 0xE3)
SAND = RGBColor(0xF3, 0xEC, 0xE0)
CREAM = RGBColor(0xFB, 0xF8, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x27, 0x34)
MUTED = RGBColor(0x5B, 0x67, 0x75)
LINE = RGBColor(0xD9, 0xD0, 0xC2)
PALE = RGBColor(0xEA, 0xF0, 0xF6)
TEAL = RGBColor(0x1F, 0x6B, 0x64)
TEAL_SOFT = RGBColor(0xE5, 0xF2, 0xF0)
GOLD = RGBColor(0xC5, 0xD0, 0xDC)

SLIDE_W = 13.333
SLIDE_H = 7.5
TOTAL = 14


def set_font(run, size: float, bold: bool, color: RGBColor) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = r_pr.find(qn(f"a:{tag}"))
        if el is None:
            el = r_pr.makeelement(qn(f"a:{tag}"), {})
            r_pr.append(el)
        el.set("typeface", FONT)


def shape(slide, kind, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, lw=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    return s


def rect(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def rrect(slide, x, y, w, h, fill: RGBColor, line: RGBColor | None = None, lw=1.0):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw)
    s.adjustments[0] = 0.08
    return s


def txt(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set(
        "anchor",
        {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
    )
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def rich(slide, blocks, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=6):
    """blocks: list of (text, size, bold, color)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.bodyPr.set(
        "anchor",
        {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
    )
    for i, (text, size, bold, color) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = text
        set_font(run, size, bold, color)
    return box


def note(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(slide, n: int, eyebrow: str, title: str, spoken: str):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    rect(slide, 0, 0, 0.12, SLIDE_H, NAVY)
    txt(slide, eyebrow, 0.48, 0.28, 10.2, 0.28, 13, True, RUST)
    txt(slide, title, 0.46, 0.56, 11.4, 0.62, 28, True, NAVY)
    txt(slide, f"{n:02d}", 12.15, 0.32, 0.85, 0.4, 16, True, RUST, PP_ALIGN.RIGHT)
    rect(slide, 0.48, 7.18, 12.4, 0.015, LINE)
    txt(slide, "禾书耕文  ·  阶段汇报  ·  2026.10.06", 0.48, 7.22, 8.2, 0.24, 11, False, MUTED)
    txt(slide, f"{n} / {TOTAL}", 10.6, 7.22, 2.25, 0.24, 11, False, MUTED, PP_ALIGN.RIGHT)
    note(slide, spoken)


def card(slide, x, y, w, h, fill=WHITE, line=LINE):
    return rrect(slide, x, y, w, h, fill, line, 1.0)


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "禾书耕文：阶段汇报"
    prs.core_properties.subject = "2026-10-06"
    prs.core_properties.author = "禾书耕文"
    cover(prs)
    toolbench(prs)
    agent(prs)
    engineering(prs)
    result(prs)
    variance(prs)
    usable(prs)
    iterate(prs)
    meeting(prs)
    assets(prs)
    later(prs)
    use_it(prs)
    budget(prs)
    report_style(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


def cover(prs) -> None:
    slide = blank(prs)
    rect(slide, 0, 0, 5.35, SLIDE_H, NAVY)
    rect(slide, 5.35, 0, SLIDE_W - 5.35, SLIDE_H, CREAM)
    rect(slide, 5.35, 0, 0.08, SLIDE_H, RUST)
    txt(slide, "阶段汇报", 0.48, 1.55, 4.5, 0.32, 14, True, RUST)
    txt(slide, "禾书耕文", 0.45, 2.05, 4.6, 0.85, 40, True, WHITE)
    txt(slide, "农业科研 AI 辅助写作", 0.48, 3.0, 4.5, 0.4, 16, False, GOLD)
    txt(slide, "2026年10月6日\n课题组", 0.48, 5.85, 4.4, 0.8, 15, False, GOLD)

    txt(slide, "今天讲三件事", 5.85, 0.55, 6.8, 0.55, 28, True, NAVY)
    items = [
        ("01", "先前的工作基础", "已经做成什么"),
        ("02", "下一步的工作方向与重点", "这一阶段重点做什么"),
        ("03", "我需要的支持", "使用、经费、汇报方式"),
    ]
    for i, (num, title, sub) in enumerate(items):
        y = 1.45 + i * 1.45
        card(slide, 5.85, y, 6.9, 1.28, WHITE, LINE)
        rect(slide, 5.85, y, 0.1, 1.28, RUST if i < 2 else NAVY)
        txt(slide, num, 6.2, y + 0.28, 1.1, 0.7, 26, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, title, 7.4, y + 0.22, 5.0, 0.48, 20, True, NAVY)
        txt(slide, sub, 7.4, y + 0.72, 5.0, 0.34, 14, False, MUTED)
    txt(slide, "讲完下一步之后：这个产品以后还可以怎么发展", 5.85, 6.85, 6.9, 0.35, 14, False, MUTED)
    note(
        slide,
        "今天按这个顺序讲。先说已经做成什么，再说下一步重点做什么，最后说我需要的支持。下一步讲完，会接着说以后还可以怎么发展。",
    )


def toolbench(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        2,
        "一  ·  先前的工作基础",
        "从往上加功能，到万能工具台",
        "接到项目时，目标就是做出能写可投稿论文的系统。当时还不清楚这是一个要长期运行、给很多人用的工程，没有先把功能模块和架构定完整，主要是到组里看大家平时要用什么，就加一块。这些功能单独都能用。叠在一起之后，新用户的上手难度变高了。",
    )
    card(slide, 0.48, 1.45, 12.35, 0.85, PALE, None)
    txt(slide, "目标", 0.72, 1.58, 1.2, 0.55, 14, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
    txt(slide, "让 AI 写出一篇可以拿去投稿的论文", 2.0, 1.58, 10.4, 0.55, 20, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)

    steps = [
        ("01", "当时", "只会做小的演示程序"),
        ("02", "做法", "看学长学姐要用什么\n就往上加一块"),
        ("03", "单独", "查重、文献、作图、写作\n都能用"),
        ("04", "叠起来", "变成万能工具台"),
    ]
    for i, (num, label, body) in enumerate(steps):
        x = 0.48 + i * 3.15
        card(slide, x, 2.55, 3.0, 2.85, WHITE, LINE)
        txt(slide, num, x + 0.22, 2.72, 2.5, 0.42, 18, True, RUST)
        txt(slide, label, x + 0.22, 3.18, 2.5, 0.4, 16, True, NAVY)
        txt(slide, body, x + 0.22, 3.7, 2.55, 1.35, 15, False, INK)

    card(slide, 0.48, 5.65, 12.35, 1.25, RUST_SOFT, None)
    txt(slide, "加到后来", 0.75, 5.82, 2.2, 0.35, 14, True, RUST)
    txt(slide, "既要知道有哪些模块，又要学怎么用", 0.75, 6.2, 11.6, 0.45, 22, True, NAVY)


def agent(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        3,
        "一  ·  先前的工作基础",
        "收成工具，前面加一个大脑",
        "方向不太对，所以改成现在这样。用户不用先记住系统里有什么。说下一步即可。",
    )
    txt(slide, "用户不用先记住系统里有什么", 0.5, 1.4, 12.2, 0.4, 16, False, MUTED)

    nodes = [
        (0.48, "用户", "对着一个对话框"),
        (3.7, "说下一步", "不用先学模块"),
        (6.92, "大脑", "大模型调用工具"),
        (10.14, "工具", "系统自己去调"),
    ]
    for i, (x, title, sub) in enumerate(nodes):
        fill = NAVY if i == 2 else WHITE
        title_c = WHITE if i == 2 else NAVY
        sub_c = GOLD if i == 2 else MUTED
        card(slide, x, 2.05, 2.85, 1.7, fill, None if i == 2 else LINE)
        txt(slide, title, x + 0.18, 2.35, 2.5, 0.5, 22, True, title_c)
        txt(slide, sub, x + 0.18, 2.95, 2.5, 0.45, 14, False, sub_c)
        if i < 3:
            txt(slide, "→", x + 2.85, 2.5, 0.4, 0.7, 22, True, RUST, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    tools = ["文献", "写作", "作图", "检查"]
    for i, name in enumerate(tools):
        x = 0.48 + i * 3.15
        card(slide, x, 4.15, 3.0, 1.15, PALE, None)
        txt(slide, name, x, 4.15, 3.0, 1.15, 22, True, NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    txt(slide, "这就是现在的 Agent", 0.5, 5.6, 12.2, 0.55, 22, True, RUST)
    txt(slide, "功能做成工具。加一个大脑，让大模型去调用这些工具。", 0.5, 6.2, 12.2, 0.4, 16, False, INK)


def engineering(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        4,
        "一  ·  先前的工作基础",
        "从演示程序到线上的工程化产品",
        "以前做的是演示用的小程序。这个产品要给多用户用，工程化是另一件难事。这些准备是为了出了问题能看见，改一处时也能知道有没有把原来的流程改坏。",
    )
    card(slide, 0.48, 1.5, 6.0, 3.35, WHITE, LINE)
    txt(slide, "以前", 0.75, 1.7, 5.4, 0.35, 14, True, MUTED)
    txt(slide, "演示用的小程序", 0.75, 2.15, 5.4, 0.5, 24, True, NAVY)
    rich(
        slide,
        [
            ("一个人：对接需求、写成代码、测试、维护", 16, False, INK),
            ("要上线，给多人用", 16, False, INK),
            ("原来的做法撑不住", 16, True, RUST),
        ],
        0.75,
        2.85,
        5.4,
        1.7,
        gap=10,
    )

    card(slide, 6.85, 1.5, 6.0, 3.35, NAVY, None)
    txt(slide, "现在", 7.12, 1.7, 5.4, 0.35, 14, True, RUST)
    txt(slide, "线上的工程化产品", 7.12, 2.15, 5.4, 0.5, 24, True, WHITE)
    txt(slide, "后面改一处，能知道有没有把原来的流程改坏。", 7.12, 2.9, 5.4, 1.3, 16, False, GOLD)

    pills = ["测试埋点", "冒烟", "管道化", "后台监测"]
    for i, name in enumerate(pills):
        x = 0.48 + i * 3.15
        card(slide, x, 5.15, 3.0, 0.85, RUST_SOFT, None)
        txt(slide, name, x, 5.15, 3.0, 0.85, 18, True, NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(slide, "现在线上这个，就是工程化做完之后的产品", 0.5, 6.25, 12.2, 0.45, 16, False, MUTED)


def result(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        5,
        "一  ·  先前的工作基础",
        "这一阶段的结果",
        "两件事一起说。工业化证明后面还能迭代。完整初稿证明产品能力，不是只有某一个功能。能跑出来，说明系统做得到。换一个人用，不一定走出同样的结果。下一部分讲这个。",
    )
    card(slide, 0.48, 1.5, 6.05, 3.15, WHITE, LINE)
    rect(slide, 0.48, 1.5, 6.05, 0.1, NAVY)
    txt(slide, "工业化产品", 0.75, 1.85, 5.5, 0.45, 22, True, NAVY)
    txt(slide, "后面还能继续改\n不用每次重做", 0.75, 2.55, 5.5, 1.5, 20, False, INK)

    card(slide, 6.8, 1.5, 6.05, 3.15, WHITE, LINE)
    rect(slide, 6.8, 1.5, 6.05, 0.1, RUST)
    txt(slide, "完整初稿", 7.08, 1.85, 5.5, 0.45, 22, True, NAVY)
    txt(slide, "从题目走到\n一整篇可导出的稿", 7.08, 2.55, 5.5, 1.5, 20, False, INK)

    card(slide, 0.48, 4.9, 12.37, 1.95, PALE, None)
    txt(slide, "大家都不是计算机出身，最直白的就是看稿", 0.75, 5.1, 11.8, 0.4, 18, True, NAVY)
    txt(slide, "这几篇是我自己用系统跑出来的。课题内容会上不展开，会后可以单独找我看。", 0.75, 5.6, 11.8, 0.85, 16, False, INK)


def variance(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        6,
        "二  ·  下一步的工作方向与重点",
        "同样的系统，不同的人，结果会差很多",
        "按原稿这个例子讲。先讲一般情况，再落到我们组。",
    )
    txt(slide, "这是一个 Agent 产品，人与人用起来差别很大", 0.5, 1.38, 12.2, 0.38, 16, False, MUTED)

    card(slide, 0.48, 1.9, 6.05, 1.85, WHITE, LINE)
    txt(slide, "纯小白", 0.75, 2.08, 5.5, 0.4, 20, True, NAVY)
    txt(slide, "同样是最顶尖的模型", 0.75, 2.58, 5.5, 0.35, 14, False, MUTED)
    txt(slide, "提示词可能完全不一样", 0.75, 2.98, 5.5, 0.45, 16, True, INK)

    card(slide, 6.8, 1.9, 6.05, 1.85, NAVY, None)
    txt(slide, "资深程序员", 7.08, 2.08, 5.5, 0.4, 20, True, WHITE)
    txt(slide, "同样是最顶尖的模型", 7.08, 2.58, 5.5, 0.35, 14, False, GOLD)
    txt(slide, "提示词可能完全不一样", 7.08, 2.98, 5.5, 0.45, 16, True, WHITE)

    rows = [
        ("认识不同", "对模型能做什么、怎么把任务拆开，质量就差一截"),
        ("链路不同", "同一套提示词，调用工具的链路也可能完全不同"),
        ("偶然性大", "按概率、根据训练数据预测下一个词"),
    ]
    for i, (k, v) in enumerate(rows):
        y = 4.0 + i * 0.95
        card(slide, 0.48, y, 12.37, 0.85, WHITE, LINE)
        txt(slide, k, 0.72, y, 2.3, 0.85, 16, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, v, 3.15, y, 9.4, 0.85, 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)


def usable(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        7,
        "二  ·  下一步的工作方向与重点",
        "对我好用，不代表对大家适用",
        "好用，是学长学姐不用先变成懂模型的人，也能走出一篇能给导师看的初稿。具体工作就是根据大家的真实反馈优化编排，减少每次乱跑。",
    )
    card(slide, 0.48, 1.48, 12.37, 1.55, WHITE, LINE)
    txt(slide, "我", 0.75, 1.65, 1.2, 0.35, 14, True, RUST)
    txt(slide, "比较熟悉这套系统，知道怎么把任务说清楚", 0.75, 2.05, 11.7, 0.4, 18, True, NAVY)
    txt(slide, "学长学姐第一次用，说法不同，工具路径就可能不同", 0.75, 2.5, 11.7, 0.35, 15, False, INK)

    card(slide, 0.48, 3.25, 5.7, 1.55, PALE, None)
    txt(slide, "上一阶段", 0.75, 3.42, 5.1, 0.32, 14, True, MUTED)
    txt(slide, "可以用", 0.75, 3.82, 5.1, 0.65, 32, True, NAVY)

    card(slide, 7.15, 3.25, 5.7, 1.55, NAVY, None)
    txt(slide, "这一阶段", 7.42, 3.42, 5.1, 0.32, 14, True, RUST)
    txt(slide, "变得好用", 7.42, 3.82, 5.1, 0.65, 32, True, WHITE)

    txt(slide, "→", 6.2, 3.55, 0.9, 0.9, 28, True, RUST, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    card(slide, 0.48, 5.05, 12.37, 1.8, RUST_SOFT, None)
    txt(slide, "按真实卡住的地方改 Agent 编排", 0.75, 5.25, 11.8, 0.45, 20, True, NAVY)
    txt(slide, "把常用写作路径收稳。不用先变成懂模型的人，也能走出一篇能给导师看的初稿。", 0.75, 5.8, 11.8, 0.7, 16, False, INK)


def iterate(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        8,
        "二  ·  下一步的工作方向与重点",
        "路很长，但已经能按反馈改",
        "工程化是为了支撑后面一轮一轮改。使用本身就是下一阶段的工作条件。",
    )
    points = [
        ("01", "不是三两天能调好", "编排要按真实使用一轮一轮收"),
        ("02", "改一处，能较快确认", "测试、管道、监测已经有了"),
        ("03", "需要大家真的用起来", "人用得少，我就还是在按自己的习惯调"),
    ]
    for i, (num, title, sub) in enumerate(points):
        y = 1.45 + i * 1.35
        card(slide, 0.48, y, 12.37, 1.22, WHITE, LINE)
        txt(slide, num, 0.75, y, 1.3, 1.22, 20, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, title, 2.15, y + 0.14, 10.2, 0.45, 20, True, NAVY)
        txt(slide, sub, 2.15, y + 0.64, 10.2, 0.4, 15, False, MUTED)

    card(slide, 0.48, 5.55, 12.37, 1.3, NAVY, None)
    txt(
        slide,
        "真实题目多了，才能按大家的习惯快速迭代",
        0.75,
        5.55,
        11.8,
        1.3,
        22,
        True,
        WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def meeting(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        9,
        "三  ·  以后还可以怎么发展",
        "组会上反复出现的一件事",
        "这是我从组会里总结出的一个大问题，也是后面想继续做的原因。",
    )
    txt(slide, "上学期一直在开组会", 0.5, 1.4, 12.2, 0.35, 15, False, MUTED)

    card(slide, 0.48, 1.95, 6.05, 2.15, WHITE, LINE)
    txt(slide, "老师常说", 0.75, 2.15, 5.5, 0.35, 14, True, RUST)
    txt(slide, "这个方向我们发过论文", 0.75, 2.6, 5.5, 1.1, 22, True, NAVY)

    card(slide, 6.8, 1.95, 6.05, 2.15, WHITE, LINE)
    txt(slide, "老师常说", 7.08, 2.15, 5.5, 0.35, 14, True, RUST)
    txt(slide, "这个实验我们做过", 7.08, 2.6, 5.5, 1.1, 22, True, NAVY)

    card(slide, 0.48, 4.4, 12.37, 2.4, NAVY, None)
    txt(slide, "人是流动的，科研是一直往下做的", 0.8, 4.7, 11.7, 0.55, 22, True, WHITE)
    txt(slide, "新来的成员不知道以前的研究基础，就会重复。", 0.8, 5.45, 11.7, 0.8, 18, False, GOLD)


def assets(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        10,
        "三  ·  以后还可以怎么发展",
        "文献留下了，其他资产也可以留下",
        "文献是同一类问题里已经在做的那一件。进展、论文、实验数据可以按同一方式留在组里。",
    )
    items = [
        ("文献", "本地知识库", "已经在做", True),
        ("研究进展", "还在各人电脑和记忆里", "可以留下", False),
        ("已发表论文", "还在各人电脑和记忆里", "可以留下", False),
        ("实验数据", "还在各人电脑和记忆里", "可以留下", False),
    ]
    for i, (name, where, tag, done) in enumerate(items):
        x = 0.48 + i * 3.15
        card(slide, x, 1.55, 3.0, 3.35, NAVY if done else WHITE, None if done else LINE)
        txt(slide, tag, x + 0.22, 1.78, 2.55, 0.35, 13, True, RUST if not done else RGBColor(0xF0, 0xC4, 0xA8))
        txt(slide, name, x + 0.22, 2.35, 2.55, 1.15, 26, True, WHITE if done else NAVY)
        txt(slide, where, x + 0.22, 3.7, 2.55, 0.85, 14, False, GOLD if done else MUTED)

    card(slide, 0.48, 5.15, 12.37, 1.7, TEAL_SOFT, None)
    txt(slide, "新同学进组，先能查到本组做过什么", 0.75, 5.4, 11.8, 0.5, 22, True, TEAL)
    txt(slide, "文献、进展、论文、实验数据，按同一方式留在组里。", 0.75, 6.0, 11.8, 0.45, 16, False, INK)


def later(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        11,
        "三  ·  以后还可以怎么发展",
        "有了这些基础之后",
        "规划方向要等这些资产都在组里之后。商业化说的是底层可以迁移，不是今天就要对外。本组先用顺，就是那套架构的一个实际样子。这也是我理解的实验室发展方向。",
    )
    tops = [
        ("01", "整理下一步", "根据已有工作，帮忙看还可以做什么"),
        ("02", "解放导师", "减少反复提醒哪些已经做过"),
        ("03", "用上组内资产", "农业人工智能，不只是帮人写一篇稿"),
    ]
    for i, (num, title, sub) in enumerate(tops):
        x = 0.48 + i * 4.2
        card(slide, x, 1.48, 4.0, 2.15, WHITE, LINE)
        txt(slide, num, x + 0.22, 1.62, 3.5, 0.32, 13, True, RUST)
        txt(slide, title, x + 0.22, 1.98, 3.55, 0.5, 20, True, NAVY)
        txt(slide, sub, x + 0.22, 2.55, 3.55, 0.8, 14, False, INK)

    card(slide, 0.48, 3.9, 12.37, 2.9, NAVY, None)
    txt(slide, "底层留下", 0.8, 4.15, 11.6, 0.35, 14, True, RUST)
    txt(slide, "写作    ·    文献    ·    作图    ·    检查", 0.8, 4.55, 11.6, 0.5, 22, True, WHITE)
    txt(slide, "换一个课题组，换上他们的文献、数据和使用习惯，就可以迁过去。", 0.8, 5.25, 11.6, 0.45, 16, False, GOLD)
    txt(slide, "本组用顺，就是以后能对外、甚至商业化的一个实际样子。", 0.8, 5.8, 11.6, 0.5, 16, True, WHITE)


def use_it(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        12,
        "四  ·  我需要的支持",
        "我希望大家用起来",
        "和成熟商业软件比，我们最贴本组的使用习惯和业务流程。大家用的时候缺什么、别扭在哪，我可以当时改。",
    )
    rows = [
        ("入口", "学习成本还在。已经收成一个对话框，说下一步即可。"),
        ("时机", "学期中还没到截止日期，写论文的需求确实不高。"),
        ("优势", "人在现场，可以按本组习惯随时加、随时改。"),
    ]
    for i, (k, v) in enumerate(rows):
        y = 1.5 + i * 1.25
        card(slide, 0.48, y, 12.37, 1.1, WHITE, LINE)
        txt(slide, k, 0.75, y, 1.6, 1.1, 18, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, v, 2.5, y, 10.0, 1.1, 18, False, INK, anchor=MSO_ANCHOR.MIDDLE)

    card(slide, 0.48, 5.4, 12.37, 1.45, RUST, None)
    txt(slide, "希望现在就拿手头正在做的内容进来用", 0.75, 5.4, 11.8, 1.45, 22, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)


def budget(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        13,
        "四  ·  我需要的支持",
        "经费",
        "四块分开说，方便分别定。阶段奖励不在今天报一个固定数。",
    )
    rows = [
        ("AI 订阅", "开发和日常使用都依赖模型", False),
        ("运行成本", "服务器这些基本开支", False),
        ("劳务", "每月 1500 元，覆盖生活费", True),
        ("阶段奖励", "看用起来之后的反馈再定，今天不定数字", False),
    ]
    for i, (k, v, hot) in enumerate(rows):
        y = 1.5 + i * 1.3
        card(slide, 0.48, y, 12.37, 1.15, NAVY if hot else WHITE, None if hot else LINE)
        txt(slide, k, 0.75, y, 2.6, 1.15, 20, True, WHITE if hot else NAVY, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, v, 3.5, y, 8.9, 1.15, 18, False, GOLD if hot else INK, anchor=MSO_ANCHOR.MIDDLE)


def report_style(prs) -> None:
    slide = blank(prs)
    chrome(
        slide,
        14,
        "四  ·  我需要的支持",
        "汇报方式",
        "我原来就有每周写周报的习惯。希望汇报还在，时间更多留在产品上。",
    )
    card(slide, 0.48, 1.5, 12.37, 1.15, RUST_SOFT, None)
    txt(slide, "每周做 PPT，会挤掉改系统的时间", 0.75, 1.5, 11.8, 1.15, 22, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)

    card(slide, 0.48, 2.9, 4.3, 3.85, WHITE, LINE)
    txt(slide, "继续留着", 0.75, 3.15, 3.8, 0.4, 14, True, MUTED)
    txt(slide, "周报", 0.75, 3.65, 3.8, 0.7, 32, True, NAVY)
    txt(slide, "文字记录继续写", 0.75, 4.5, 3.8, 1.2, 16, False, INK)

    card(slide, 5.05, 2.9, 7.8, 3.85, NAVY, None)
    txt(slide, "会上口头三句", 5.35, 3.15, 7.2, 0.4, 14, True, RUST)
    lines = ["这周做了什么", "大家卡在哪", "下周改什么"]
    for i, line in enumerate(lines):
        txt(slide, f"0{i + 1}", 5.35, 3.7 + i * 0.85, 0.7, 0.6, 18, True, RUST, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, line, 6.15, 3.7 + i * 0.85, 6.2, 0.6, 22, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)


if __name__ == "__main__":
    build()
